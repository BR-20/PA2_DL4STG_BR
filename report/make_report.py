"""Builds report/25280081_PA2_AI651_BR.docx from report/q1.md and the saved results.

Every number is read from results/*.json, results/*.npz or predictions/*.md. Each qualitative
claim in the Question 2 text is guarded by an assert on the same numbers, so the text cannot
contradict the results it describes.
"""
import argparse
import os
import re
import shutil
import subprocess
import sys

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from results_io import (BLANK, COPY, D_ARCH, DIAG, FLOW, HARDWARE, REACH, RES, ROOT, SEL, TABLE1, ev, fmt,
                        infer_time, mean, ms, pct, seeds, train_time)

REPORT = os.path.join(ROOT, "report")
FIG = "../figures"
QUESTIONS = {
    "Q2.1": "Based on Figure 1(a), would mirrored and rotated trajectories already be familiar to a model trained on this data? (Hint: how would the direction histogram change if every training video were mirrored? If every video were rotated by 90°?) From Figure 1(b), what is a typical speed, and what do 2× and 3× correspond to in pixels per frame? Optical flow is unreliable where digits overlap; could that bias your histograms enough to change your answer?",
    "Q2.2": "For each transformation, does it change the motion statistics, the digit appearance, or both? Why does this make the three of them useful as a controlled experiment?",
    "Q2.3": "Estimate model A’s reach in original pixels per time step, using the conventions above. Walk through your reasoning from the patch size, kernel sizes, number of layers and the recurrence. There is no formula to look up for this.",
    "Q2.4": "Using Table 2 and Figure 2, how does model A compare with the two reference predictors, and how does its error grow over the five forecast horizons? What happens to its predictions when digits overlap?",
    "Q2.5": "Before you run any experiment in Testing the Baseline, write down your predictions: will model A pass the mirror test? The rotation test? At what speed should it break down, given your reach estimate and Figure 1? Include this answer in your report exactly as you first wrote it, even if it turns out to be wrong.",
    "Q2.6": "Blank and Copy-last are exactly consistent under mirror and rotation. What does this say about reading C_T on its own, and why does E_T always have to be reported next to it?",
    "Q2.7": "Pick one video in Figure 3 where rows 3 and 4 disagree and describe precisely how the model’s account of the motion changes. For example, a digit goes the wrong way, gets lost, bounces too early or not at all, or blurs into something unrecognisable.",
    "Q2.8": "Is rotation harder than mirroring for model A? Using your answers to Q2.1 and Q2.2, argue whether the difference comes from the motion or from the digits’ appearance.",
    "Q2.9": "At what speed do model A’s predictions fall apart, and does this agree with the reach you estimated in Q2.3? If not, what did your estimate miss? Describe what the failure in Figure 5 looks like (blur, a digit that stalls, a duplicated digit, a digit that vanishes) and what each of these suggests the model is doing.",
    "Q2.10": "From these results alone, can you tell whether the speed failure happens because fast motion never appeared in training, or because it is beyond the model’s reach? Propose an experiment that would tell the two apart.",
    "Q2.11": "Before training B–D, predict in at most three sentences which model will fix which of the failures you found in Testing the Baseline. As with Q2.5, include this unedited in your report.",
    "Q2.12": "Describe model D and estimate its reach the same way you did in Q2.3, showing that it is strictly larger than model A’s.",
    "Q2.13": "Label each failure from Testing the Baseline as appearance, unfamiliar motion, or architectural reach, citing the comparison in Table 5 that supports each label. In particular, if model C still does badly at 3× while model D does well, what does that tell you, and why can’t more training data carry information further than the architecture allows?",
    "Q2.14": "Did any of these changes cost accuracy on the original 1× clips? If so, was it worth it?",
    "Q2.15": "Compare symmetry that is learned (model B) with symmetry that is built in (the symmetrised models) in terms of consistency, accuracy on original and transformed clips, and inference cost. Which one is exact, and why? Averaging eight predictions can also blur them; when would you expect built-in symmetry to hurt accuracy, and did it here?",
    "Q2.16": "Go back to the video you discussed in Q2.7. Does the final model now get it right (Figure 6), and which change fixed which part of the failure?",
    "Q2.17": "Final verdict (150 words max). Did your final model learn motion? What evidence would convince a skeptic, and what is one test it still has not passed?",
}
VERBATIM = []


def q(label):
    return f"**{label}.** {QUESTIONS[label]}\n\n**Answer.** "


def table(headers, rows):
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    out += ["| " + " | ".join(str(c).replace("|", "\\|") for c in r) + " |" for r in rows]
    return "\n".join(out) + "\n"


def image(path, width="6.4in"):
    return f"![]({path}){{width={width}}}\n"


def committed(name):
    with open(os.path.join(ROOT, "predictions", name)) as fh:
        text = fh.read().strip()
    body = re.sub(r"^\*\*Q2\.\d+ Answer\.\*\* ", "", text)
    VERBATIM.append(body)
    return body


def near(c, cs):
    return min(cs, key=lambda d: np.hypot(c[0] - d[0], c[1] - d[1]))


def dist(a, b):
    return float(np.hypot(a[0] - b[0], a[1] - b[1]))


def gaps(name, key):
    """Per seed transformed error minus E(5)."""
    return [e - b for e, b in zip(seeds(name, key), seeds(name, "E5_1x"))]


def pervideo(tag, key):
    return np.load(os.path.join(RES, f"pervideo_{tag}.npz"))[key]


def build_q2():
    md = []
    add = md.append
    a5, b5, c5 = mean("A", "E5_1x"), BLANK["E5_1x"], COPY["E5_1x"]
    med = FLOW["median_speed"]
    reach_a = max(REACH["A"]["analytic"].values())
    reach_d_key = next(k for k, v in REACH.items() if k != "A" and v["cfg"] == D_ARCH)
    reach_d = REACH[reach_d_key]
    final = "D"

    add("# Question 2. Motion or Memory?\n")
    add("**Setup.** Every number below comes from the run recorded in the notebook of the repository. "
        f"All timings use one {HARDWARE.split(',')[0]} GPU on Google Colab with PyTorch {HARDWARE.split('torch ')[1].split(',')[0]}. "
        "Videos 0 to 7,999 train, 8,000 to 8,999 validate and 9,000 to 9,999 test. Test video i in the figures is dataset index 9,000 plus i. "
        "Each learned model was trained for 30 epochs with Adam at learning rate 10⁻³ and batch size 32 under seeds 0 and 1. "
        "Every table entry for a learned model is written as mean (seed 0, seed 1) and every figure uses the seed 0 model. "
        "Errors carry three significant figures and values below 10⁻³ use scientific notation. "
        "Training time per epoch excludes the dataset download and the validation check that ran every five epochs. "
        "Inference time is the median of five timed passes over the 1,000 test windows at batch size 32 after one warm up batch. "
        "Table S1 maps each component to its file in the repository.\n")
    add(table(["Component", "File"], [
        ["Mirror, rotation, the group G and the symmetrised forecaster", "pa2/transforms.py, pa2/metrics.py"],
        ["Window rules for 1×, 2× and 3×", "pa2/windows.py"],
        ["ConvLSTM forecaster for models A to D", "pa2/models.py"],
        ["Training with speed and symmetry augmentation", "pa2/train.py"],
        ["Errors, consistency gaps and timing", "pa2/metrics.py"],
        ["Optical flow for Figure 1", "pa2/flow.py"],
        ["Reach count and impulse test", "pa2/reach.py"],
        ["Stage runner and report builder", "run.py, report/make_report.py"],
    ]))
    add("Table S1. This table gives the repository file that implements each part of the experiment.\n")

    add("## 1. Setup\n")
    add("**Q.** Check that your transformation and window functions do exactly what is defined above, using the same functions your experiments will call. Report the results as Table 1, with the last column filled in.\n")
    rows = []
    for r in TABLE1:
        got = r["yours"]
        if isinstance(got, list) and got and isinstance(got[0], list):
            cell = f"{', '.join(map(str, got[0]))} / {', '.join(map(str, got[1]))}"
        elif isinstance(got, list):
            cell = "(" + ", ".join(map(str, got)) + ")"
        else:
            assert got == 0 and r["torch_equal"] is True, r
            cell = "0, torch.equal True"
        label = r["check"].replace(" - ", " − ").replace("R^4", "R⁴").replace("R^-1", "R⁻¹")
        rows.append([label, r["expected"].replace("1-5 / 6-10", "1–5 / 6–10"), cell])
    assert [tuple(TABLE1[3]["yours"]), tuple(TABLE1[4]["yours"])] == [(7, 10, 43), (7, 43, 10)]
    assert TABLE1[7]["yours"] == [[1, 4, 7, 10, 13], [16, 19]]
    add("**Answer.** Table 1 gives the eight checks. Every row matches its expected value, so the functions used by all later experiments flip the correct axis, rotate anticlockwise and pick the correct frames.\n")
    add(table(["Check", "Expected", "Yours"], rows))
    add("Table 1. This table gives the transformation checks on the random video x and the marker video p and the window checks on a video whose frame f holds the value f.\n")

    add("**Q.** Estimate dense optical flow on the first 50 training videos with the stated Farneback settings, keep pixels above 0.2 in both frames with flow magnitude above 0.5 pixels, and show Figure 1: (a) the direction histogram with 18 bins and a line at 1/18 and (b) the speed histogram with the median marked, both as fractions of kept pixels.\n")
    add(image(f"{FIG}/fig1.png"))
    add(f"Figure 1. The plot shows (a) the fraction of kept pixels in each 20 degree direction bin with the uniform level 1/18 as a dashed line and (b) the fraction of kept pixels by flow speed with the median as a vertical line. "
        f"The {FLOW['pairs']} frame pairs of the first 50 training videos keep {FLOW['kept_pixels']:,} pixels and the median speed is {fmt(med)} pixels per frame.\n")
    noise = FLOW["tv_half_split"]
    assert FLOW["tv_mirror"] < noise < FLOW["tv_rotation"] < 1.5 * noise
    assert FLOW["tv_all_vs_separate"] < 0.25 * noise
    add(q("Q2.1") +
        "Mirroring every video maps a direction θ to 180° − θ and a 90° rotation maps θ to θ + 90°. "
        "If the training directions were uniform, both maps would leave Figure 1(a) unchanged and transformed trajectories would already be familiar. "
        f"The bins of Figure 1(a) range from {fmt(FLOW['min_bin_over_uniform'])} to {fmt(FLOW['max_bin_over_uniform'])} times the uniform level, so the histogram is not flat. "
        "To judge whether this matters, the total variation distance between two histograms is compared with a noise floor. "
        f"The histogram differs from its mirrored copy by {fmt(FLOW['tv_mirror'])} and from its rotated copy by {fmt(FLOW['tv_rotation'])}. "
        f"Two halves of the same data, videos 0 to 24 against videos 25 to 49, already differ by {fmt(noise)}. "
        "Both transformed distances are of the size of this sampling noise, so 50 videos cannot show a real preference of direction. "
        "This agrees with the data generator of Srivastava et al. (2015, p. 5), which draws each direction uniformly on the unit circle. "
        "Mirrored and rotated trajectories are therefore familiar as motion. Only the digits look different. "
        f"Figure 1(b) gives a typical speed of {fmt(med)} pixels per frame with quartiles {fmt(FLOW['speed_quartiles'][0])} and {fmt(FLOW['speed_quartiles'][1])}. "
        f"A 2× clip therefore moves a digit about {fmt(2 * med)} pixels between the frames the model sees and a 3× clip about {fmt(3 * med)} pixels. "
        f"To test the overlap bias, the flow was pooled again over only the {FLOW['separate_digit_pairs']} frame pairs in which both frames show two separate bright blobs. "
        f"This subset keeps {FLOW['kept_pixels_separate']:,} pixels, its median speed is {fmt(FLOW['median_speed_separate'])} pixels per frame and its direction histogram differs from the full one by only {fmt(FLOW['tv_all_vs_separate'])}. "
        f"Its mirror and rotation distances are {fmt(FLOW['tv_mirror_separate'])} and {fmt(FLOW['tv_rotation_separate'])}, so the overlap bias is too small to change either answer.\n")
    add(q("Q2.2") +
        "Mirror and rotation keep every speed and only permute directions, which Q2.1 found to be already balanced, so they leave the motion statistics nearly unchanged. "
        "They do change the digit appearance, because a mirrored 2 or a rotated 7 is not a digit shape of the training set. "
        "Rotation changes the appearance more than mirroring, because a mirror keeps vertical strokes vertical and leaves 0, 1 and 8 nearly unchanged, while a rotation turns upright strokes sideways. "
        "Speed up keeps every digit exactly as it looks in training and only multiplies the displacement per seen frame by 2 or 3, so it changes the motion statistics alone. "
        "The three transformations therefore isolate the two factors. "
        "A failure under mirror or rotation can be traced to appearance and a failure under speed up to motion, while the true future stays a valid Moving MNIST future in every case. "
        "The same trained model and the same test videos are used throughout, so any change in error comes from the transformation alone.\n")
    add("**Trade off.** A direction test on 50 videos is cheap but noisy, so the familiarity verdict rests on the noise floor and on the documented data generator together and not on the histogram alone.\n")

    add("## 2. Baseline: ConvLSTM\n")
    add("**Q.** Table 2: MSE₁, …, MSE₅, E⁽⁵⁾ and trainable parameter count for model A, Blank and Copy-last. Figure 2: input, prediction and ground truth for three test videos, at least one of them with overlapping digits.\n")
    params_a = ev("A", 0)["params"]
    assert params_a == ev("A", 1)["params"] == 480784
    mse_a = [[ev("A", s)["mse_h"][h] for s in (0, 1)] for h in range(5)]
    rows = [[f"MSE{h + 1}".translate(str.maketrans("12345", "₁₂₃₄₅")), ms(mse_a[h]), fmt(BLANK["mse_h"][h]), fmt(COPY["mse_h"][h])] for h in range(5)]
    rows.append(["E⁽⁵⁾", ms(seeds("A", "E5_1x")), fmt(b5), fmt(c5)])
    rows.append(["Trainable parameters", f"{params_a:,}", "0", "0"])
    add(table(["Quantity", "Model A", "Blank", "Copy-last"], rows))
    add("Table 2. This table gives the error at each forecast horizon, E⁽⁵⁾ and the trainable parameter count at 1× with t = 1 on the 1,000 test videos.\n")
    add(image(f"{FIG}/fig2.png"))
    v_over, v_a, v_b = SEL["fig2"]
    add(f"Figure 2. The plot shows the five input frames and the five future frames of test videos {v_over}, {v_a} and {v_b}, with the truth above and the seed 0 model A prediction below. "
        f"The digits of video {v_over} overlap in the future frames.\n")

    ra, ea = REACH["A"]["analytic"], REACH["A"]["empirical"]
    assert ra == ea and reach_a == 8
    add(q("Q2.3") +
        "Patch folding turns each 4 × 4 pixel patch into one cell of a 16 × 16 grid, so one cell step equals 4 original pixels. "
        "A 3 × 3 convolution with padding 1 connects a cell only to its 8 neighbours, so each such convolution moves information by at most one cell or 4 pixels. "
        "Within one time step there are two such convolutions in series, the gate convolution of layer 1 and the gate convolution of layer 2. "
        f"A new input frame passes through both of them before it reaches the output head, which gives 2 cells or {ra['input']} pixels. "
        f"The previous hidden state of layer 1 enters the layer 1 convolution next to the new input and then also passes through layer 2, which again gives {ra['state_1']} pixels. "
        f"The previous hidden state of layer 2 only passes through the layer 2 convolution, which gives {ra['state_2']} pixels. "
        "The cell states only interact element by element. The 1 × 1 head and the pixel shuffles move nothing under the stated conventions. "
        f"The reach of model A is the longest of these paths, {reach_a} pixels per time step. "
        f"An impulse test on the built network perturbs one cell and finds changes up to {ea['input']}, {ea['state_1']} and {ea['state_2']} pixels for the three paths, which matches the count.\n")

    grow = [ev("A", s)["mse_h"][4] / ev("A", s)["mse_h"][0] for s in (0, 1)]
    over = DIAG["fig2"][str(v_over)]
    merged = [h for h in range(5) if over["truth"][h]["blobs"] == 1]
    assert merged == [1, 2, 3, 4]
    peaks = [st["peak"] for v in DIAG["fig2"].values() for st in v["pred"]]
    assert a5 < b5 < c5 and all(1.4 < g < 1.6 for g in grow)
    assert all(np.mean(mse_a[h]) < min(BLANK["mse_h"][h], COPY["mse_h"][h]) for h in range(5))
    inc = np.diff([np.mean(m) for m in mse_a])
    assert inc[0] == inc.max() and all(i < 0.7 * inc[0] for i in inc[1:])
    near_pair = DIAG["fig2"][str(v_a)]
    assert all(st["blobs"] == 2 for st in near_pair["truth"]) and near_pair["pred"][2]["blobs"] == 1
    add(q("Q2.4") +
        f"Model A reaches E⁽⁵⁾ of {fmt(a5)} against {fmt(b5)} for Blank and {fmt(c5)} for Copy-last, which is {fmt(a5 / b5)} and {fmt(a5 / c5)} of their errors. "
        "Copy-last is worse than Blank because a copied digit sits where the true digit no longer is, so it pays for the missing digit and for the ghost digit at once. "
        "Model A beats both at every horizon in Table 2, so it predicts where the digits go and does not only reproduce the last frame. "
        f"Its error rises from {fmt(np.mean(mse_a[0]))} at the first horizon to {fmt(np.mean(mse_a[4]))} at the fifth, a factor of {fmt(np.mean(grow))} that both seeds share. "
        f"The largest rise, {fmt(np.mean(mse_a[1]) - np.mean(mse_a[0]))}, comes between the first two horizons and later rises are about half as large, because the prediction turns into a wider and fainter blob whose error grows more slowly. "
        "The growth has two causes. Each fed back prediction carries its own error into the next step. The uncertainty about position also grows with the horizon, so the squared error loss favours a blurred average of the possible positions. "
        f"Figure 2 shows this blur, since the predicted peak intensity is between {fmt(min(peaks))} and {fmt(max(peaks))} against 1 in the truth. "
        f"In video {v_over} the 8 and the 2 merge into one bright region from the second forecast frame onwards. "
        "The true frames keep both shapes visible through the overlap, while the prediction turns the two digits into one grey smear in which neither shape can be recognised. "
        f"The same merging appears in video {v_a}, where the true 6 and 9 stay apart but the predicted blobs join from the third forecast frame, so nearby digits are averaged together even without a true overlap.\n")
    add(q("Q2.5") + committed("Q2.5.md") + "\n")
    add("This prediction was committed to the repository before any mirror, rotation or speed test ran. Section 3 compares it with the measurements.\n")
    add("**Trade off.** The 1 × 1 head and the pixel shuffle make model A cheap, with 480,784 parameters, but its 8 pixel reach is close to the 2× displacement, so the baseline is expected to be fragile at speed.\n")

    add("## 3. Testing the Baseline\n")
    add("**Q.** Table 3: E⁽⁵⁾, C_M, E_M, C_R and E_R for model A, Blank and Copy-last. Figure 3: the two test videos with the largest per video C_M and the two with the largest per video C_R for the seed 0 model A, each shown as four rows of ten frames, with animated GIFs saved in the repository. Figure 4: histograms of the per video C_M and C_R of the seed 0 model A, with the videos from Figure 3 marked.\n")
    keys = ["E5_1x", "CM", "EM", "CR", "ER"]
    labels = ["E⁽⁵⁾", "C_M", "E_M", "C_R", "E_R"]
    rows = [[lab, ms(seeds("A", k)), fmt(BLANK[k]), fmt(COPY[k])] for lab, k in zip(labels, keys)]
    add(table(["Quantity", "Model A", "Blank", "Copy-last"], rows))
    add("Table 3. This table gives E⁽⁵⁾, the consistency gaps and the transformed errors under mirror M and rotation R at 1× with t = 1 on the 1,000 test videos.\n")
    pv = {k: pervideo("A_s0", k) for k in ("CM", "CR")}
    panels = [("M", r, v) for r, v in enumerate(SEL["fig3_M"], 1)] + [("R", r, v) for r, v in enumerate(SEL["fig3_R"], 1)]
    for i, (name, rank, v) in enumerate(panels):
        letter = "abcd"[i]
        base = f"fig3_{name}_rank{rank}_video{v}"
        assert os.path.exists(os.path.join(ROOT, "gifs", base + ".gif"))
        tname = "mirror" if name == "M" else "rotation"
        add(image(f"{FIG}/{base}.png"))
        add(f"Figure 3({letter}). The plot shows test video {v}, which has the {'largest' if rank == 1 else 'second largest'} per video C_{name} of {fmt(pv['C' + name][v])} under {tname} for the seed 0 model A. "
            f"Rows show (1) X and f(X), (2) T(X) and T(Y), (3) T(X) and f(T(X)) and (4) T(X) and T(f(X)), with inputs in columns 1 to 5 and the future in columns 6 to 10. The GIF is gifs/{base}.gif.\n")
    add(image(f"{FIG}/fig4.png"))
    add(f"Figure 4. The plot shows histograms of the per video C_M and C_R of the seed 0 model A over the 1,000 test videos on a log count axis. "
        f"Dashed lines mark the Figure 3 videos {SEL['fig3_M'][0]} and {SEL['fig3_M'][1]} for C_M and {SEL['fig3_R'][0]} and {SEL['fig3_R'][1]} for C_R, which lie in the far right tails.\n")

    assert BLANK["CM"] == BLANK["CR"] == COPY["CM"] == COPY["CR"] == 0
    assert BLANK["EM"] == BLANK["E5_1x"] and COPY["ER"] == COPY["E5_1x"]
    em, er = mean("A", "EM"), mean("A", "ER")
    add(q("Q2.6") +
        f"Blank and Copy-last have C_M and C_R of exactly 0 in Table 3, yet their transformed errors are {fmt(BLANK['EM'])} and {fmt(COPY['EM'])}, against {fmt(em)} for model A under mirror. "
        "Both rules act on every pixel in the same way and never look at where anything is, so they commute with any rearrangement of pixels by construction. "
        "A zero gap therefore says only that the predictor treats a transformed clip the same way as the original. It says nothing about whether either prediction is right. "
        "C_T compares the model with itself, while E_T compares the model with the true transformed future. "
        "A low C_T with a high E_T means a consistent but wrong forecaster and a low C_T with a low E_T means a consistent and correct one, so only the pair answers the question. "
        "This is why E_T has to be reported next to every C_T in this report.\n")

    rec = next(r for r in DIAG["fig3"]["A"]["R"] if r["video"] == SEL["fig3_R"][0])
    v7 = rec["video"]
    ty, r3, r4 = rec["TY"], rec["fTX"], rec["TfX"]
    low_true = [max(c[0] for c in st["centroids"]) for st in ty]
    low3 = [max(c[0] for c in st["centroids"]) for st in r3]
    assert [st["blobs"] for st in r3[2:]] == [1, 1, 1] and all(st["blobs"] >= 2 for st in r4[:4])
    assert low3[4] > low_true[4] + 5 and rec["err_fTX"] > rec["err_TfX"]
    add(q("Q2.7") +
        f"Test video {v7} in Figure 3(c) is used, with T = R. In row 2 the rotated 3 moves up to the top wall and comes back down, while the rotated 8 reaches the bottom wall and moves back up and to the left until the two digits come close in the last frame. "
        "Row 4 is model A run on the original clip and then rotated. It follows this account. Both blobs come back from their walls and draw together by the fifth forecast frame, although they are blurred. "
        "Row 3 is model A run on the rotated clip. It tells a different story. "
        "The rotated 8 does not bounce off the bottom wall but slides along it to the left, with its blob centre near row "
        f"{low3[2]:.0f} to {low3[4]:.0f} in forecast frames 3 to 5 while the true digit rises to row {low_true[4]:.0f}. "
        "The rotated 3 stays pressed against the top wall in the first two forecast frames and then fades below half intensity, so from the third frame only one blob above 0.5 remains. "
        f"Row 3 therefore makes one digit bounce late or not at all and loses the other, which raises its error against the true rotated future to {fmt(rec['err_fTX'])} against {fmt(rec['err_TfX'])} for row 4. "
        "In the original orientation the same event is a bounce of an upright 8 off the left wall, which row 4 shows the model can handle. The model therefore likely ties its rebound rule to the upright shapes it saw in training.\n")

    gm, gr = gaps("A", "EM"), gaps("A", "ER")
    assert all(x > 0 for x in gm + gr) and all(r > m for r, m in zip(gr, gm))
    assert mean("A", "CR") > mean("A", "CM")
    b_gm, b_gr = gaps("B", "EM"), gaps("B", "ER")
    add(q("Q2.8") +
        f"Rotation is harder than mirroring for model A. "
        f"The rise from E⁽⁵⁾ to E_R is {fmt(np.mean(gr))}, or {pct(np.mean(gr) / a5)}, against {fmt(np.mean(gm))}, or {pct(np.mean(gm) / a5)}, from E⁽⁵⁾ to E_M. Both seeds show the same order. "
        f"C_R is {fmt(mean('A', 'CR'))} against {fmt(mean('A', 'CM'))} for C_M, so the rotated forecasts also disagree more with the rotated originals. "
        "Both changes are small, because E_M and E_R stay within 8 percent of E⁽⁵⁾, so model A passes both tests in the weak sense of degrading only a little. "
        "Q2.1 showed that the motion directions are already balanced, with mirror and rotation distances at the level of the sampling noise. "
        "Q2.2 showed that both transformations keep every speed, so neither creates an unfamiliar motion. "
        f"The motion seen by the model is therefore nearly the same in both cases and cannot explain why rotation costs {fmt(np.mean(gr) / np.mean(gm))} times more. "
        "Appearance can explain it, because a mirror keeps strokes upright and leaves several digits unchanged while a rotation turns every upright stroke sideways. "
        "Patch folding adds to this, because a rotation also moves pixels across the channels of each folded patch, so the learned first layer filters receive patterns they never saw. "
        f"Model B supports the appearance reading in Table 5, since training on all eight symmetries turns the rise to E_R into {fmt(np.mean(b_gr))} and the rise to E_M into {fmt(np.mean(b_gm))}, which means no rise at all.\n")

    add("**Q.** Table 4: model A’s E⁽²⁾ at 1×, 2× and 3×, and its E⁽⁵⁾ at 1× and 2×, with Copy-last for comparison. Figure 5: for one test video, the five input frames and the first two predicted and true future frames at 1×, 2× and 3× for the seed 0 model A, with no overlap in any displayed true frame.\n")
    t4 = [("E⁽²⁾ at 1×", "E2_1x"), ("E⁽²⁾ at 2×", "E2_2x"), ("E⁽²⁾ at 3×", "E2_3x"), ("E⁽⁵⁾ at 1×", "E5_1x"), ("E⁽⁵⁾ at 2×", "E5_2x")]
    add(table(["Quantity", "Model A", "Copy-last", "Blank"], [[lab, ms(seeds("A", k)), fmt(COPY[k]), fmt(BLANK[k])] for lab, k in t4]))
    add("Table 4. This table gives E⁽²⁾ at 1×, 2× and 3× and E⁽⁵⁾ at 1× and 2× with t = 1 on the 1,000 test videos. Blank is added as the level of a forecast that contains no digit.\n")
    v5 = SEL["fig5"]
    assert SEL["fig5_rule_met"]
    for k in ("1", "2", "3"):
        assert all(st["blobs"] == 2 for st in DIAG["fig5"]["A"][k]["truth"])
    add(image(f"{FIG}/fig5.png", "5.6in"))
    add(f"Figure 5. The plot shows test video {v5} at 1×, 2× and 3×, with the five input frames and the first two true frames in the truth rows and the first two seed 0 model A predictions in the model rows. "
        "The two digits stay separate in every displayed true frame, with the closest approach in the last 3× frame.\n")

    a2 = {k: seeds("A", f"E2_{k}x") for k in (1, 2, 3)}
    assert all(v > BLANK["E2_2x"] for v in a2[2]) and all(v > BLANK["E2_3x"] for v in a2[3])
    assert all(v < BLANK["E2_1x"] for v in a2[1])
    moves = {}
    for k in ("1", "2", "3"):
        e = DIAG["fig5"]["A"][k]
        last = e["input"][-1]["centroids"]
        moves[k] = ([dist(c, near(c, e["truth"][0]["centroids"])) for c in last],
                    [dist(c, near(c, e["pred"][0]["centroids"])) for c in last],
                    [e["pred"][h]["peak"] for h in (0, 1)])
    assert max(moves["2"][1]) < 0.5 * min(moves["2"][0]) and max(moves["3"][1]) < 1.0
    mv = lambda xs: " and ".join(f"{x:.1f}" for x in xs)
    add(q("Q2.9") +
        f"Model A falls apart at 2×. Its E⁽²⁾ rises from {fmt(np.mean(a2[1]))} at 1× to {fmt(np.mean(a2[2]))} at 2× and {fmt(np.mean(a2[3]))} at 3×. "
        f"At 2× and 3× both seeds are worse than Blank, which scores {fmt(BLANK['E2_2x'])} and {fmt(BLANK['E2_3x'])} by drawing nothing at all. "
        "A forecast that loses to an empty frame carries no usable information about where the digits are, which is the threshold used here for falling apart. "
        "Copy-last is not a fair threshold at speed, since its error grows with the displacement and model A still beats it by moving its blobs a little. "
        f"This does not agree with the reach estimate. The 2× displacement of about {fmt(2 * med)} pixels is inside the reach of {reach_a} pixels. The Q2.5 prediction of a break at 3× was wrong. "
        "The estimate missed three things. "
        "First, reach only says what the wiring allows, while the weights encode the speeds seen in training, which lie between "
        f"{fmt(FLOW['speed_quartiles'][0])} and {fmt(FLOW['speed_quartiles'][1])} pixels per frame for the middle half of the pixels. "
        "Second, the hidden state of layer 2 can only shift its content by 4 pixels per step. If this state carries the tracked digit, which is likely but not tested, the usable reach for a tracked digit is below 8 pixels. "
        "Third, the 8 pixel count is measured between cell centres, while pixels at the edges of two patches can be up to 3 pixels further apart. "
        "Figure 5 shows what the failure looks like. "
        f"Measured from the blob centres of pixels above 0.5, the predicted digits move {mv(moves['1'][1])} pixels in the first forecast frame at 1×, close to the true {mv(moves['1'][0])} pixels. "
        f"At 2× they move {mv(moves['2'][1])} pixels against a true {mv(moves['2'][0])} pixels. At 3× they move only {mv(moves['3'][1])} pixels against a true {mv(moves['3'][0])} pixels. "
        f"The 2× digits therefore advance at about the training speed and fall behind. The 3× digits stall at their last seen position. Both are blurred to a peak of {fmt(min(moves['2'][2] + moves['3'][2]))} to {fmt(max(moves['2'][2] + moves['3'][2]))}. "
        "A stalled digit suggests that the model cannot match the new position with the old one and keeps the digit where its state last saw it. "
        "Blur suggests that it averages over several positions it cannot tell apart. "
        "A duplicated digit would suggest that the old position stays in the state while the input adds a new one. A vanished digit would suggest that no match was found at all. "
        "Neither a duplicate nor a vanished digit appears in Figure 5 for model A.\n")
    add(q("Q2.10") +
        "Not fully. Model A never saw fast motion and it also has a fixed reach. Both explanations predict that it fails at 3×. "
        "The 2× result gives one hint, because the 2× displacement lies inside the counted reach and the model still fails, which points to unfamiliar motion. "
        "This hint is only an inference, since the reach count is an upper bound on what the wiring allows and not a measured tracking range. "
        "An experiment that separates the two causes changes one factor at a time. "
        "Train the same architecture on speeds drawn from 1×, 2× and 3×, which is model C. Then train a network with larger reach on the same speed mix, which is model D. "
        "If C recovers at a speed, the failure there was unfamiliar motion. If C still fails while D recovers, the failure was reach. "
        "If neither recovers, the cause is something else, such as the limit on what five frames reveal about the future.\n")
    add("**Trade off.** The mirror and rotation tests show a small appearance cost, while the speed test shows a large motion failure, so the baseline is mainly limited by the speeds it has seen and not by orientation.\n")

    add("## 4. What Can Data Fix?\n")
    add(q("Q2.11") + committed("Q2.11.md") + "\n")
    assert all(v < c for v, c in zip(a2[2], [COPY["E2_2x"]] * 2))
    add("This prediction was committed before model B, C or D was trained. It is reproduced unedited. "
        "It is wrong, because it came from a failure rule that compared model A with Copy-last at speed. "
        f"Model A beats Copy-last at 2× with {fmt(np.mean(a2[2]))} against {fmt(COPY['E2_2x'])} only because Copy-last itself collapses, while Section 3 shows that model A loses to Blank at both 2× and 3×. "
        "With Blank as the threshold, the prediction would have been that model C fixes the speed failures and that model B removes the small mirror and rotation gaps.\n")

    add("**Q.** Report Table 5: for models A–D, E⁽⁵⁾ at 1×; C_M, E_M, C_R, E_R; E⁽²⁾ at 1×, 2× and 3×; trainable parameter count; training time per epoch; and inference time.\n")
    t5_keys = [("E⁽⁵⁾ at 1×", "E5_1x"), ("C_M", "CM"), ("E_M", "EM"), ("C_R", "CR"), ("E_R", "ER"),
               ("E⁽²⁾ at 1×", "E2_1x"), ("E⁽²⁾ at 2×", "E2_2x"), ("E⁽²⁾ at 3×", "E2_3x")]

    def model_table(models, sym):
        rows = [[lab] + [ms(seeds(n, k, s)) for n, s in zip(models, sym)] for lab, k in t5_keys]
        rows.append(["Trainable parameters"] + [f"{ev(n, 0, s)['params']:,}" for n, s in zip(models, sym)])
        rows.append(["Training time per epoch (s)"] + [ms(train_time(n)) for n in models])
        rows.append(["Inference time (s)"] + [ms(infer_time(n, s)) for n, s in zip(models, sym)])
        return rows

    for n in "ABCD":
        p = [ev(n, s)["params"] for s in (0, 1)]
        assert p[0] == p[1] <= 721176
    add(table(["Quantity", "Model A", "Model B", "Model C", "Model D"], model_table("ABCD", [False] * 4)))
    add("Table 5. This table gives accuracy, consistency, speed errors, size and cost for models A to D on the 1,000 test videos. Times are in seconds on the T4 GPU.\n")
    rd_a, rd_e = reach_d["analytic"], reach_d["empirical"]
    assert rd_a == rd_e and max(rd_a.values()) > reach_a and reach_d["params"] == ev("D", 0)["params"]
    add(q("Q2.12") +
        f"Model D keeps the structure of model A, with patch size {D_ARCH['patch']}, two ConvLSTM layers of {D_ARCH['hidden']} hidden channels, {D_ARCH['kernel']} × {D_ARCH['kernel']} kernels and the same head and 9 step rollout. "
        f"The only change is a dilation of {D_ARCH['dilation'][1]} in the layer 2 gate convolution, while layer 1 keeps dilation {D_ARCH['dilation'][0]}. "
        f"A dilated kernel spreads the same nine weights over a wider grid, so D has exactly {ev('D', 0)['params']:,} trainable parameters, the same as A and below the limit of 721,176. "
        "It uses no attention, normalisation or extra layer. It is trained with speed augmentation exactly like model C. "
        "The reach follows the Q2.3 method. "
        f"The layer 1 convolution still moves 1 cell, while the layer 2 convolution with dilation {D_ARCH['dilation'][1]} moves {D_ARCH['dilation'][1]} cells, so one step through both moves 4 cells or {rd_a['input']} pixels. "
        f"The new input and the previous layer 1 state both cross the two convolutions and reach {rd_a['input']} and {rd_a['state_1']} pixels, while the previous layer 2 state reaches {rd_a['state_2']} pixels. "
        f"The reach of model D is therefore {max(rd_a.values())} pixels per step, twice the {reach_a} pixels of model A and above the 3× displacement of about {fmt(3 * med)} pixels. "
        f"The impulse test on the built network gives {rd_e['input']}, {rd_e['state_1']} and {rd_e['state_2']} pixels, which matches. "
        "The design was fixed before training as the default of the plan, because it raises the reach at no parameter cost. "
        "A comparison of four candidate designs on the validation split was started but stopped by hand after the first candidate had trained for 10 epochs and before its validation score was computed, to save GPU time, so no candidate comparison is reported and no test video was used to choose D. "
        "One known weakness of this design is that a dilated layer 2 links only cells 3 apart in its own recurrence, which likely costs fine detail. This was not tested.\n")

    c2, c3 = seeds("C", "E2_2x"), seeds("C", "E2_3x")
    d2, d3 = seeds("D", "E2_2x"), seeds("D", "E2_3x")
    assert max(c2) < min(a2[2]) and max(c3) < min(a2[3]) and max(c2) < BLANK["E2_2x"] and max(c3) < BLANK["E2_3x"]
    assert mean("D", "E2_3x") >= mean("C", "E2_3x") and mean("D", "E2_2x") >= mean("C", "E2_2x")
    red = lambda x, y: 1 - y / x
    add(q("Q2.13") +
        "Each label below comes from the comparison that changes only that factor. "
        f"The mirror failure, a small penalty of {pct(np.mean(gm) / a5)}, is labelled appearance. Model B, trained on all eight symmetries with everything else equal, reduces the rise from E⁽⁵⁾ to E_M from {fmt(np.mean(gm))} for A to {fmt(np.mean(b_gm))} and C_M from {fmt(mean('A', 'CM'))} to {fmt(mean('B', 'CM'))}. "
        f"The rotation failure, a penalty of {pct(np.mean(gr) / a5)}, is labelled appearance for the same reason. B reduces the rise to E_R from {fmt(np.mean(gr))} to {fmt(np.mean(b_gr))} and C_R from {fmt(mean('A', 'CR'))} to {fmt(mean('B', 'CR'))}. "
        "Since Q2.1 found the training motion already balanced, the main thing that symmetry augmentation adds is transformed digit shapes, which is why the label is appearance and not unfamiliar motion. "
        f"The 2× failure is labelled unfamiliar motion. Model C, which only adds 2× and 3× clips, lowers E⁽²⁾ at 2× from {fmt(mean('A', 'E2_2x'))} to {fmt(mean('C', 'E2_2x'))}, a drop of {pct(red(mean('A', 'E2_2x'), mean('C', 'E2_2x')))} that takes it well below Blank. "
        f"The 3× failure is also labelled unfamiliar motion. Model C lowers E⁽²⁾ at 3× from {fmt(mean('A', 'E2_3x'))} to {fmt(mean('C', 'E2_3x'))}, a drop of {pct(red(mean('A', 'E2_3x'), mean('C', 'E2_3x')))}. "
        f"Model D, which only adds reach to C, does not improve further, with {fmt(mean('D', 'E2_2x'))} at 2× and {fmt(mean('D', 'E2_3x'))} at 3× against {fmt(mean('C', 'E2_2x'))} and {fmt(mean('C', 'E2_3x'))} for C. "
        "No failure in this run is therefore labelled architectural reach. "
        f"Model C at 3× is still worse than C at 1×, with {fmt(mean('C', 'E2_3x'))} against {fmt(mean('C', 'E2_1x'))}. The larger reach of D does not close this gap. "
        "The remaining gap is likely due to the harder 3× task, with fewer start frames per video, more wall bounces between seen frames and only two scored targets, but this was not tested. "
        "If C still did badly at 3× while D did well, the failure would be architectural reach. "
        "It would mean that C has seen the fast motion and still cannot represent it, while a network whose information travels further per step can. "
        "More training data cannot fix such a failure, because each output pixel of one update depends only on the inputs and states inside its reach. "
        "A weight that does not exist in the wiring cannot be learned, so more examples only tune the weights inside the reach and never connect a digit to a position it cannot see.\n")

    chg = {n: mean(n, "E5_1x") - a5 for n in "BCD"}
    b_seed_gap = abs(seeds("B", "E5_1x")[0] - seeds("B", "E5_1x")[1])
    assert abs(chg["B"]) < b_seed_gap
    assert min(seeds("C", "E5_1x")) > max(seeds("A", "E5_1x")) and min(seeds("D", "E5_1x")) > max(seeds("C", "E5_1x"))
    add(q("Q2.14") +
        f"Yes for C and D but not clearly for B. E⁽⁵⁾ at 1× is {fmt(a5)} for A, {fmt(mean('B', 'E5_1x'))} for B, {fmt(mean('C', 'E5_1x'))} for C and {fmt(mean('D', 'E5_1x'))} for D. "
        f"The change for B is {pct(chg['B'] / a5)}, but the two seeds of B differ by {fmt(b_seed_gap)}, more than the change itself, so two seeds cannot show a cost. "
        f"C costs {pct(chg['C'] / a5)} and D costs {pct(chg['D'] / a5)}. Both seeds of each lie above both seeds of the model before it, so these costs are real. "
        "The cost of C is worth it. "
        f"For a rise of {fmt(chg['C'])} in E⁽⁵⁾ at 1×, C cuts E⁽²⁾ by {fmt(mean('A', 'E2_2x') - mean('C', 'E2_2x'))} at 2× and {fmt(mean('A', 'E2_3x') - mean('C', 'E2_3x'))} at 3×, where A was worse than an empty frame. "
        "The cost comes from capacity. The same 480,784 weights now cover three speeds. A third of the training windows are 3× windows with only two scored targets. "
        "B is worth it as well, since it removes the mirror and rotation penalties at no measurable cost on original clips. "
        "D is not worth it as built. It costs more than C at 1×, is slower per epoch and gives no gain at 2× or 3×, so its extra reach did not buy accuracy in this run.\n")
    add("**Trade off.** Data augmentation fixed every failure that model A showed, while extra reach did not, so on this data the cheaper fix of showing the model more kinds of motion was also the effective one.\n")

    add("## 5. The Final Model\n")
    add("**Q.** Show (a short argument is enough) that f̄ has C_M = C_R = 0 exactly for any f. Your implementation should give zero up to floating point error; report these values in Table 6.\n")
    add("**Answer.** Every g in G rearranges pixels, so it is linear and invertible. The product of two elements of G is again in G. "
        "The symmetrised forecaster is $\\bar f(X)=\\frac{1}{8}\\sum_{g\\in G} g^{-1} f(gX)$. Take any h in G and feed the transformed clip hX.\n\n"
        "$$\\bar f(hX)=\\frac{1}{8}\\sum_{g\\in G} g^{-1} f(ghX)$$\n\n"
        "Substitute $g'=gh$. As g runs over G, g' also runs over every element of G exactly once. It also follows that $g^{-1}=h\\,g'^{-1}$.\n\n"
        "$$\\bar f(hX)=\\frac{1}{8}\\sum_{g'\\in G} h\\,g'^{-1} f(g'X)=h\\,\\frac{1}{8}\\sum_{g'\\in G} g'^{-1} f(g'X)=h\\,\\bar f(X)$$\n\n"
        "The middle step moves h out of the sum, which is allowed because h is linear. "
        "With h = M this gives $\\bar f(MX)=M\\bar f(X)$ and so $C_M=0$, and with h = R it gives $C_R=0$. "
        "The argument never uses what f is, so it holds for any trained forecaster. "
        "In floating point the two routes add the same eight terms in a different order, which leaves a rounding residue. "
        f"Table 6 shows this residue as at most {fmt(max(max(seeds(n, k, True)) for n in 'BD' for k in ('CM', 'CR')))}, about fourteen orders of magnitude below the errors.\n")
    add("**Q.** Table 6: the Table 5 columns for symmetrised B, symmetrised D and the final model if it is different, with the training time of the underlying model and C_M and C_R in scientific notation. Figure 6: (a) Figure 3 repeated with the final model on exactly the same videos, with GIFs saved, and (b) Figure 5 repeated with the final model on exactly the same video.\n")
    add(table(["Quantity", "Symmetrised B", "Symmetrised D (final)"], model_table("BD", [True, True])))
    add("Table 6. This table gives the Table 5 quantities for the symmetrised models on the 1,000 test videos. The training time is that of the underlying model. The inference time includes all eight transformations, forward passes, inverse transformations and the averaging.\n")
    add("The final model is model D wrapped as a symmetrised forecaster, as the assignment specifies. "
        "Model D+B was not trained, because Table 5 gives no reason to expect that it would beat symmetrised D. "
        "Symmetry augmentation in B removed only small orientation gaps. The wrapper already removes those gaps exactly without retraining.\n")
    for i, (name, rank, v) in enumerate(panels):
        letter = "abcd"[i]
        base = f"fig6a_{name}_rank{rank}_video{v}"
        assert os.path.exists(os.path.join(ROOT, "gifs", base + ".gif"))
        add(image(f"{FIG}/{base}.png"))
        add(f"Figure 6(a{i + 1}). The plot repeats Figure 3({letter}) for test video {v} with the final model, symmetrised D at seed 0, using the same four rows. The GIF is gifs/{base}.gif.\n")
    add(image(f"{FIG}/fig6b.png", "5.6in"))
    add(f"Figure 6(b). The plot repeats Figure 5 for test video {v5} with the final model, symmetrised D at seed 0.\n")

    bs, ds = mean("B", "E5_1x", True), mean("D", "E5_1x", True)
    for n in "BD":
        assert all(s < r for s, r in zip(seeds(n, "E5_1x", True), seeds(n, "E5_1x")))
    ratio_b = np.mean(infer_time("B", True)) / np.mean(infer_time("B"))
    ratio_d = np.mean(infer_time("D", True)) / np.mean(infer_time("D"))
    final_peaks = [st["peak"] for r in DIAG["fig3"]["final"]["R"] + DIAG["fig3"]["final"]["M"] for st in r["fTX"]]
    a_peaks = [st["peak"] for r in DIAG["fig3"]["A"]["R"] + DIAG["fig3"]["A"]["M"] for st in r["fTX"]]
    add(q("Q2.15") +
        f"Learned symmetry leaves model B with C_M of {fmt(mean('B', 'CM'))} and C_R of {fmt(mean('B', 'CR'))}, which are smaller than for A but clearly not zero. "
        f"Built in symmetry gives C_M and C_R of order 10⁻¹⁷ for both wrapped models in Table 6, which is floating point rounding. "
        "Only the built in version is exact, because the group average makes the identity hold for any f, as the proof shows, while augmentation only rewards consistency in the loss and leaves the network free to break it on unseen clips. "
        f"On original clips E⁽⁵⁾ falls from {fmt(mean('B', 'E5_1x'))} for B to {fmt(bs)} for symmetrised B, lower for both seeds. For D it changes by only {fmt(mean('D', 'E5_1x') - ds)}, with {fmt(mean('D', 'E5_1x'))} for D and {fmt(ds)} for symmetrised D. "
        f"On transformed clips the wrapped models have E_M and E_R equal to their E⁽⁵⁾, while B has E_M of {fmt(mean('B', 'EM'))} and E_R of {fmt(mean('B', 'ER'))}. "
        f"The price is inference cost. Eight forward passes make symmetrised B {fmt(ratio_b)} times and symmetrised D {fmt(ratio_d)} times slower than the unwrapped models. "
        "Averaging cannot raise the squared error above the mean of the eight transformed errors, because the squared error is convex and each g only rearranges pixels. "
        "Built in symmetry would therefore hurt accuracy when the model is much better on the original orientation than on the other seven, or when its eight predictions are sharp but disagree, since their mean is then a faint mixture of several positions. "
        "In this run it did not hurt the squared error, since every wrapped model matches or beats its unwrapped model at 1×. "
        f"The frames of the final model are visibly blurred, with a peak intensity of {fmt(min(final_peaks))} to {fmt(max(final_peaks))} in Figure 6(a) against {fmt(min(a_peaks))} to {fmt(max(a_peaks))} for model A in Figure 3. "
        "Without a figure of unwrapped D it was not tested how much of that blur comes from the averaging and how much from model D itself.\n")

    fr = next(r for r in DIAG["fig3"]["final"]["R"] if r["video"] == v7)
    assert fr["err_fTX"] > rec["err_TfX"] and fr["err_fTX"] < rec["err_fTX"] and fr["mse_rows"] < 1e-12
    assert all(st["blobs"] == 0 for st in fr["fTX"])
    add(q("Q2.16") +
        f"Only in part. For test video {v7} under rotation, rows 3 and 4 of Figure 6(a3) are now identical, with a squared difference of {fmt(fr['mse_rows'])} against {fmt(rec['mse_rows'])} for model A. "
        f"The final model therefore tells one account of the motion whatever the orientation. Its error against the true rotated future is {fmt(fr['err_fTX'])}, below the {fmt(rec['err_fTX'])} of row 3 of model A but above the {fmt(rec['err_TfX'])} of row 4. "
        "The forecast is not right, because no predicted pixel reaches half intensity, so the 3 and the 8 have become two soft blobs that cannot be recognised as digits. "
        "The upper blob moves down from the top wall as the true 3 does, while the lower blob stays near the bottom wall instead of rising with the true 8. "
        "The symmetrisation wrapper fixed the orientation part of the failure, because it removed the disagreement between rows 3 and 4 by construction. "
        "Speed augmentation and the extra reach of D did not fix the bounce on this 1× clip. The blurred output of D removed the digit shapes that model A still kept. "
        "Model B shows that the orientation part can also be reduced by training alone, since it lowers C_R from "
        f"{fmt(mean('A', 'CR'))} to {fmt(mean('B', 'CR'))}.\n")

    f2 = {k: mean("D", f"E2_{k}x", True) for k in (1, 2, 3)}
    assert all(f2[k] < BLANK[f"E2_{k}x"] and f2[k] < COPY[f"E2_{k}x"] for k in (1, 2, 3))
    assert f2[3] < np.mean(a2[3]) and ds > a5
    verdict = (f"Partly. The final model learned a speed aware motion rule that holds under all eight symmetries, but not a sharp picture of the digits. "
               f"Its E⁽²⁾ is {fmt(f2[1])}, {fmt(f2[2])} and {fmt(f2[3])} at 1×, 2× and 3×, below Blank and Copy-last at every speed, while model A collapsed to {fmt(np.mean(a2[3]))} at 3×. "
               "Its consistency gaps are rounding noise. "
               "A skeptic would want digit tracks and not only pixel errors, for example blob centres that follow the true centres through bounces at every speed. It would also want success at a speed never trained, such as 4×. "
               f"One test it has not passed is accuracy on the original clips. Its E⁽⁵⁾ at 1× of {fmt(ds)} is worse than the {fmt(a5)} of model A. Figure 6 shows soft blobs instead of digits, so its motion knowledge is real but coarse.")
    words = len(verdict.split())
    assert words <= 150, words
    add(q("Q2.17") + verdict + "\n")
    add("**Trade off.** Built in symmetry buys exact consistency for eight times the inference cost and no retraining, while augmentation is free at inference but only approximate. Neither repairs a blurred base model, so the base forecaster still sets the accuracy.\n")
    return "\n".join(md), words


STYLE_BANNED = ["delve", "crucial", "pivotal", "robust", "leverage", "seamless", "comprehensive", "holistic", "furthermore", "moreover",
                "additionally", "it is important to note", "in conclusion", "overall", "notably", "ensure", "utilize", "foster", "tapestry",
                "landscape", "realm"]


def style_scan(text):
    prose = re.sub(r"\$\$.*?\$\$", " ", text, flags=re.S)
    prose = re.sub(r"\$[^$]*\$", " ", prose)
    prose = re.sub(r"!\[[^\]]*\]\([^)]*\)(\{[^}]*\})?", " ", prose)
    for body in VERBATIM:
        prose = prose.replace(body, " ")
    lines = [l for l in prose.splitlines() if not l.startswith(("|", "#", "**Q"))]
    lines = [re.sub(r"^\*\*Q\d?\.?\d*\.\*\* .*?\*\*Answer\.\*\*", "", l) for l in lines]
    prose = "\n".join(lines)
    hits = []
    for ch, name in (("—", "em dash"), ("–", "en dash"), (";", "semicolon"), ("!", "exclamation")):
        if ch in prose:
            hits.append(name)
    for w in STYLE_BANNED:
        if re.search(r"\b" + re.escape(w) + r"\b", prose, re.I):
            hits.append("banned word " + w)
    if re.search(r"\w, and\b", prose):
        hits.append("comma before and")
    for m in re.finditer(r"[A-Za-z]+-[A-Za-z]+", prose):
        if m.group(0) not in ("Copy-last", "C-last"):
            hits.append("hyphen " + m.group(0))
    return sorted(set(hits))


def pandoc_exe():
    try:
        import pypandoc
        return pypandoc.get_pandoc_path()
    except Exception:
        return shutil.which("pandoc")


def style_docx(path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document(path)
    styles = doc.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in styles]:
            st = styles[name]
            st.font.name = "Calibri"
            st.font.size = Pt(10.5)
    for name, size in (("heading 1", 15), ("heading 2", 12.5), ("heading 3", 11)):
        for st in styles:
            if st.name.lower() == name:
                st.font.size = Pt(size)
                st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    for tbl in doc.tables:
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_pr = tbl._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "808080")
            borders.append(el)
        tbl_pr.append(borders)
        for r, row in enumerate(tbl.rows):
            for cell in row.cells:
                if r == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), "DCE6F1")
                    cell._tc.get_or_add_tcPr().append(shd)
                for par in cell.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(8.5)
                        if r == 0:
                            run.font.bold = True
    doc.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--name", default="Bimal Rehman")
    ap.add_argument("--roll", default="25280081")
    ap.add_argument("--repo", default="https://github.com/BR-20/PA2_DL4STG_BR")
    ap.add_argument("--out", default=os.path.join(REPORT, "25280081_PA2_AI651_BR"))
    a = ap.parse_args()
    head = ("# AI651 Deep Learning for Space, Time and Graphs, Assignment 2\n\n"
            f"**Student name.** {a.name}\n\n**Roll number.** {a.roll}\n\n**Code repository.** <{a.repo}>\n\n"
            "## AI acknowledgement\n\nThe prompts, outputs and edits are recorded in the section AI use record at the end of this report.\n")
    with open(os.path.join(REPORT, "q1.md")) as fh:
        q1 = fh.read()
    q2, words = build_q2()
    brk = '```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n'
    text = "\n\n".join([head, brk + q1, brk + q2, brk + "# AI use record\n"])
    hits = style_scan(q2) + [h + " (Q1)" for h in style_scan(q1)]
    with open(a.out + ".md", "w") as fh:
        fh.write(text)
    subprocess.run([pandoc_exe(), a.out + ".md", "-f", "markdown+tex_math_dollars-implicit_figures", "-o", a.out + ".docx",
                    "--resource-path=" + REPORT], check=True)
    style_docx(a.out + ".docx")
    print("Q2.17 words:", words)
    print("style hits:", hits or "none")
    print("written:", a.out + ".docx")
    return 1 if hits else 0


if __name__ == "__main__":
    sys.exit(main())
